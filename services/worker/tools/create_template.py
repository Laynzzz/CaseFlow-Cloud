"""Generate a synthetic, plain-placeholder DOCX fixture; no external content."""
from pathlib import Path
from docx import Document

target = Path(__file__).resolve().parents[3] / "infrastructure/local/generated/purchase-template.docx"
target.parent.mkdir(parents=True, exist_ok=True)
document = Document()
document.add_heading("Approved purchase request", 0)
document.add_paragraph("Synthetic CaseFlow training record")
for label, key in [
    ("Case", "case_id"), ("Vendor", "vendor"), ("Purchase", "description"),
    ("Currency", "currency"), ("Total", "total"), ("Cost center", "cost_center"),
    ("Business justification", "justification"), ("Line items", "line_items"),
    ("Approved at", "approved_at"), ("Reviewers", "approvers"),
]:
    document.add_heading(label, 2)
    document.add_paragraph("{{ " + key + " }}")
document.save(target)
print("Created synthetic purchase template in ignored local fixtures.")
