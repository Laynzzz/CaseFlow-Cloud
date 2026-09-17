"""Actual child-process termination around production transaction boundaries."""
from hashlib import sha256
from io import BytesIO
import json
import os
import time
from uuid import uuid4

from docx import Document
import psycopg
from psycopg.types.json import Jsonb
import pytest

from caseflow_worker import jobs, render
from caseflow_worker.settings import BUCKET, database, storage
from process_probe import terminate_at_boundary


def selected_state():
    with database() as db:
        return {
            "job": db.execute("SELECT * FROM worker.jobs").fetchone(),
            "inbox": db.execute("SELECT count(*) AS n FROM worker.inbox").fetchone()["n"],
            "artifacts": db.execute("SELECT * FROM worker.artifacts").fetchall(),
            "success_events": db.execute(
                "SELECT count(*) AS n FROM worker.outbox WHERE payload->>'status'='SUCCEEDED'"
            ).fetchone()["n"],
        }


def reclaim_after_real_expiry(old):
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        current = jobs.claim(uuid4())
        if current:
            assert current["job_id"] == old["job_id"]
            assert current["fence"] > old["fence"]
            assert current["attempt"] == old["attempt"]
            return current
        time.sleep(0.05)
    pytest.fail("crashed lease was not reclaimable before the deadline")


@pytest.mark.parametrize("boundary,committed", [("schedule-before-commit", False), ("schedule-after-commit", True)])
def test_scheduling_survives_process_termination(event, boundary, committed, record_testsuite_property):
    evidence = terminate_at_boundary(boundary, event.model_dump_json())
    record_testsuite_property(boundary, json.dumps(evidence, default=str))
    state = selected_state()
    assert state["inbox"] == int(committed)
    assert (state["job"] is not None) == committed
    assert jobs.schedule(event) is not committed
    state = selected_state()
    assert state["inbox"] == 1 and state["job"]["status"] == "QUEUED"
    assert not jobs.schedule(event)


@pytest.mark.parametrize("boundary,committed", [("result-before-commit", False), ("result-after-commit", True)])
def test_result_selection_survives_process_termination(event, boundary, committed, record_testsuite_property):
    evidence = terminate_at_boundary(boundary, event.model_dump_json())
    record_testsuite_property(boundary, json.dumps(evidence, default=str))
    old, artifact = evidence["job"], evidence["artifact"]
    state = selected_state()
    assert state["job"]["status"] == ("SUCCEEDED" if committed else "RUNNING")
    assert len(state["artifacts"]) == int(committed)
    assert state["success_events"] == int(committed)
    if committed:
        assert not jobs.schedule(event)
        assert jobs.claim(uuid4()) is None
    else:
        current = reclaim_after_real_expiry(old)
        assert not jobs.heartbeat(old)
        assert not jobs.finish(old, artifact)
        assert not jobs.fail(old, "STALE_PROCESS", True)
        assert jobs.finish(current, dict(artifact, key=artifact["key"] + "-recovered"))
    assert not jobs.finish(old, artifact)
    assert not jobs.fail(old, "DELAYED_FAILURE", True)
    state = selected_state()
    assert state["job"]["status"] == "SUCCEEDED"
    assert len(state["artifacts"]) == 1 and state["success_events"] == 1


@pytest.fixture
def uploaded_document(event, isolated_database):
    assert os.getenv("S3_ENDPOINT", "http://127.0.0.1:8333") == "http://127.0.0.1:8333"
    assert BUCKET == "caseflow-local"
    client = storage()
    job_id = uuid4()
    prefix = f"tenants/{event.tenantId}/documents/{job_id}/"
    template_key = f"tests/process-crash/{uuid4()}/template.docx"
    template = Document()
    template.add_paragraph("Vendor: {{ vendor }}")
    template.add_paragraph("Total: {{ currency }} {{ total }}")
    template.add_paragraph("{{ line_items }}")
    output = BytesIO()
    template.save(output)
    source = output.getvalue()
    snapshot = dict(
        caseId=str(event.aggregateId), approvedAt="2026-09-17T00:00:00Z",
        approvers=[{"name": "Synthetic finance"}],
        template={"key": template_key, "sha256": sha256(source).hexdigest()},
        purchase=dict(vendor="Synthetic crash-test supplies", description="Desk lamps", currency="USD",
                      total="70.00", costCenter="OPS-TEST", justification="Synthetic recovery fixture",
                      lineItems=[dict(description="Desk lamp", quantity="2", unitPrice="35.00")]),
    )
    digest = sha256(json.dumps(snapshot, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
    request = event.model_copy(update={"eventId": uuid4(), "jobId": job_id, "correlationId": job_id, "inputHash": digest})
    try:
        client.put_object(Bucket=BUCKET, Key=template_key, Body=source, IfNoneMatch="*")
        with psycopg.connect(**isolated_database) as db:
            actor = db.execute("SELECT requested_by FROM core.job_requests WHERE job_id=%s", (event.jobId,)).fetchone()[0]
            db.execute("""INSERT INTO core.job_requests(tenant_id,job_id,case_id,kind,input,input_hash,requested_by)
                VALUES (%s,%s,%s,'DOCUMENT',%s,%s,%s)""",
                       (event.tenantId, job_id, event.aggregateId, Jsonb(snapshot), digest, actor))
        yield request, client
    finally:
        # This random tenant/job belongs only to this disposable test fixture.
        for page in client.get_paginator("list_objects_v2").paginate(Bucket=BUCKET, Prefix=prefix):
            for item in page.get("Contents", []):
                assert item["Key"].startswith(prefix)
                client.delete_object(Bucket=BUCKET, Key=item["Key"])
        client.delete_object(Bucket=BUCKET, Key=template_key)


def object_bytes(client, key):
    with client.get_object(Bucket=BUCKET, Key=key)["Body"] as body:
        return body.read()


def test_uploaded_object_survives_crash_without_selecting_stale_output(uploaded_document, record_testsuite_property):
    event, client = uploaded_document
    evidence = terminate_at_boundary("after-object-upload", event.model_dump_json())
    record_testsuite_property("after-object-upload", json.dumps(evidence, default=str))
    old, orphan = evidence["job"], evidence["artifact"]
    original = object_bytes(client, orphan["key"])
    assert sha256(original).hexdigest() == orphan["sha256"]
    state = selected_state()
    assert state["job"]["status"] == "RUNNING"
    assert state["artifacts"] == [] and state["success_events"] == 0
    current = reclaim_after_real_expiry(old)
    selected = render.execute(current, jobs.input_for(current))
    assert selected["key"] != orphan["key"]
    assert jobs.finish(current, selected)
    assert not jobs.finish(old, orphan)
    assert not jobs.fail(old, "LATE_FAILURE", True)
    saved = selected_state()
    assert saved["job"]["status"] == "SUCCEEDED"
    assert len(saved["artifacts"]) == 1 and saved["success_events"] == 1
    assert saved["artifacts"][0]["object_key"] == selected["key"]
    assert object_bytes(client, orphan["key"]) == original
    downloaded = object_bytes(client, selected["key"])
    assert sha256(downloaded).hexdigest() == saved["artifacts"][0]["sha256"]
    text = "\n".join(p.text for p in Document(BytesIO(downloaded)).paragraphs)
    assert "Synthetic crash-test supplies" in text and "USD 70.00" in text and "{{" not in text


@pytest.mark.parametrize("setting,value", [
    ("DB_NAME", "caseflow"), ("DB_NAME", "caseflow_test_not-a-uuid"),
    ("DB_HOST", "remote.invalid"), ("DB_PORT", "5432"),
])
def test_process_probe_refuses_non_fixture_database(monkeypatch, setting, value):
    monkeypatch.setenv("DB_NAME", "caseflow_test_" + uuid4().hex)
    monkeypatch.setenv("DB_HOST", "127.0.0.1")
    monkeypatch.setenv("DB_PORT", "54320")
    monkeypatch.setenv(setting, value)
    with pytest.raises(ValueError, match="process probe requires"):
        terminate_at_boundary("schedule-before-commit", "{}")
