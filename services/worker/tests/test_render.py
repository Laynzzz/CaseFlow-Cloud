from io import BytesIO
from docx import Document
from caseflow_worker.render import render_bytes


def test_approved_snapshot_is_rendered_with_xml_escaping():
    template=Document()
    template.add_paragraph("Vendor: {{ vendor }}")
    template.add_paragraph("Total: {{ currency }} {{ total }}")
    template.add_paragraph("{{ line_items }}")
    source=BytesIO();template.save(source)
    snapshot=dict(caseId="synthetic",approvedAt="2026-09-15T00:00:00Z",approvers=[{"name":"Finance"}],
                  purchase=dict(vendor="Synthetic <Tools> & Co",description="Laptop",currency="USD",total="4200.00",
                                costCenter="TRAINING",justification="Synthetic test",lineItems=[dict(description="Laptop",quantity="1",unitPrice="4200.00")]))
    result=Document(BytesIO(render_bytes(source.getvalue(),snapshot)))
    text="\n".join(p.text for p in result.paragraphs)
    assert "Synthetic <Tools> & Co" in text
    assert "USD 4200.00" in text
    assert "Laptop: 1 × 4200.00" in text
    assert "{{" not in text
